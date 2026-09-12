import json
from pathlib import Path

from docx import Document

from app.utils.logging import get_logger

logger = get_logger("document_loader")


def _load_docx(docx_path: Path) -> str:
    """Extract paragraph and table text from a .docx file."""
    doc = Document(str(docx_path))
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text_parts.append("\t".join(cells))
    return "\n".join(text_parts)


def load_docx_files(directory: str) -> list[dict]:
    """Load every .docx file in a directory into raw text documents."""
    documents = []
    docx_dir = Path(directory)
    if not docx_dir.exists():
        logger.warning("DOCX directory %s does not exist", directory)
        return documents

    for docx_path in sorted(docx_dir.glob("*.docx")):
        if docx_path.name.startswith("~$"):
            continue
        logger.info("Loading %s", docx_path.name)
        try:
            text = _load_docx(docx_path)
            if text.strip():
                documents.append(
                    {
                        "page_content": text,
                        "metadata": {"source": docx_path.name},
                    }
                )
        except Exception as exc:
            logger.error("Failed to parse %s: %s", docx_path.name, exc)

    return documents


def _sanitize_metadata(entry: dict) -> dict:
    """Keep only Chroma-compatible scalar metadata (drop None/dicts, join lists)."""
    meta = {}
    for key, value in entry.items():
        if key == "content_chunk":
            continue
        if value is None:
            continue
        if isinstance(value, (list, tuple)):
            value = ", ".join(str(item) for item in value)
        if isinstance(value, (str, int, float, bool)):
            meta[key] = value
    return meta


def _entry_lists(data) -> list[list[dict]]:
    """Find the list-of-dicts arrays in a JSON document (tables/registries)."""
    known_keys = ("directorio", "directorio_contactos", "cuentas_bancarias", "registros")
    for key in known_keys:
        value = data.get(key)
        if isinstance(value, list) and value and isinstance(value[0], dict):
            return [value]
    lists = [
        value
        for value in data.values()
        if isinstance(value, list) and value and isinstance(value[0], dict)
    ]
    return lists if lists else [data.get("directorio", [])]


def _build_content_chunk(entry: dict) -> str:
    """Generate readable text from a JSON row when no content_chunk is provided."""
    lines = []
    for key, value in entry.items():
        if value is None:
            continue
        if isinstance(value, (list, tuple)):
            value = ", ".join(str(item) for item in value)
        if not isinstance(value, (str, int, float, bool)):
            continue
        label = key.replace("_", " ").title()
        lines.append(f"{label}: {value}")
    return "\n".join(lines)


def load_json_files(directory: str) -> list[dict]:
    """Load structured JSON entries as prechunked documents."""
    documents = []
    json_dir = Path(directory)
    if not json_dir.exists():
        return documents

    for json_path in sorted(json_dir.glob("*.json")):
        logger.info("Loading %s", json_path.name)
        try:
            data = json.loads(json_path.read_text(encoding="utf-8"))
            for entries in _entry_lists(data):
                for entry in entries:
                    if not isinstance(entry, dict):
                        continue
                    content = entry.get("content_chunk") or _build_content_chunk(entry)
                    if not content:
                        logger.warning("Empty entry in %s", json_path.name)
                        continue
                    documents.append(
                        {
                            "page_content": content,
                            "metadata": {
                                "source": json_path.name,
                                "prechunked": True,
                                **_sanitize_metadata(entry),
                            },
                        }
                    )
        except Exception as exc:
            logger.error("Failed to parse %s: %s", json_path.name, exc)

    return documents

def load_markdown_files(directory: str) -> list[dict]:
    """Load every .md file in a directory into raw text documents."""
    documents = []
    md_dir = Path(directory)
    if not md_dir.exists():
        return documents

    for md_path in sorted(md_dir.glob("*.md")):
        logger.info("Loading %s", md_path.name)
        try:
            text = md_path.read_text(encoding="utf-8")
            if text.strip():
                documents.append(
                    {
                        "page_content": text,
                        "metadata": {
                            "source": md_path.name,
                            "format": "markdown",
                        },
                    }
                )
        except Exception as exc:
            logger.error("Failed to parse %s: %s", md_path.name, exc)

    return documents


def load_documents(
    docx_dir: str, md_dir: str | None = None, json_dir: str | None = None
) -> list[dict]:
    """Load DOCX, Markdown, and JSON documents from the given directories."""
    return (
        load_docx_files(docx_dir)
        + load_markdown_files(md_dir or docx_dir)
        + load_json_files(json_dir or docx_dir)
    )
